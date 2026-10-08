#!/usr/bin/env python3
"""Regenerate public packet PDFs.

The repo previously stored these PDFs as static binaries (python-docx →
LibreOffice) with no generator. This script renders the updated markdown
sources with fpdf2 and Inter, the face already embedded in the prior files.

Outputs:
  public/docs/DABonData.pdf          ← public/docs/DABonData.md
  public/docs/DABonData.docx         ← same markdown, so the docx is not stale
  public/docs/DAB_Energy_Trust_Memo.pdf
  public/docs/DAB_Energy_Trust_Bylaws.pdf
  public/docs/DAB_Resident_Proposal_Letter.pdf
  public/docs/DAB_Script_Handout.pdf
  public/docs/DAB_Example_Email.pdf
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "public" / "docs"
SRC = ROOT / "docs" / "pdf-src"
FONT = Path("/usr/share/fonts/truetype/macos")
NAVY = (23, 54, 93)
SECTION_BLUE = (79, 129, 189)
INK = (20, 16, 12)


def asterisks_to_fpdf(text: str) -> str:
    """fpdf2 markdown italics use __, not *single asterisks*."""
    bolds: list[str] = []

    def stash(match: re.Match[str]) -> str:
        bolds.append(match.group(0))
        return f"\x00B{len(bolds) - 1}\x00"

    text = re.sub(r"\*\*[^*]+\*\*", stash, text)
    text = re.sub(r"\*([^*\n]+)\*", r"__\1__", text)

    def restore(match: re.Match[str]) -> str:
        return bolds[int(match.group(1))]

    return re.sub(r"\x00B(\d+)\x00", restore, text)


RUN_RE = re.compile(r"\*\*([^*]+)\*\*|\*([^*\n]+)\*")


def add_md_runs(paragraph, text: str, *, force_bold: bool = False) -> None:
    pos = 0
    for match in RUN_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            if force_bold:
                run.bold = True
        if match.group(1) is not None:
            run = paragraph.add_run(match.group(1))
            run.bold = True
        else:
            run = paragraph.add_run(match.group(2))
            run.italic = True
            if force_bold:
                run.bold = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        if force_bold:
            run.bold = True

PACKET = (DOCS / "DABonData.md").read_text(encoding="utf-8")


def slice_between(text: str, start: str, end: str | None) -> str:
    i = text.find(start)
    if i < 0:
        raise SystemExit(f"missing start marker: {start}")
    if end is None:
        return text[i:].strip() + "\n"
    j = text.find(end, i + len(start))
    if j < 0:
        raise SystemExit(f"missing end marker: {end}")
    return text[i:j].strip() + "\n"


class PacketPDF(FPDF):
    def __init__(self) -> None:
        super().__init__(format="letter", unit="mm")
        self.add_font("Inter", "", str(FONT / "Inter-Regular.ttf"))
        self.add_font("Inter", "B", str(FONT / "Inter-Bold.ttf"))
        self.add_font("Inter", "I", str(FONT / "Inter-Italic.ttf"))
        self.add_font("Inter", "BI", str(FONT / "Inter-BoldItalic.ttf"))
        self.set_auto_page_break(auto=True, margin=18)
        self.set_margins(16, 16, 16)

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Inter", "", 8)
        self.set_text_color(90, 78, 60)
        self.cell(
            0,
            6,
            f"DAB ENERGY TRUST — proposed, not enacted law  |  {self.page_no()}",
            align="C",
        )
        self.set_text_color(20, 16, 12)

    def heading(self, text: str, level: int) -> None:
        text = asterisks_to_fpdf(text)
        plain = re.sub(r"[*_]", "", text).strip()
        self.ln(3 if level > 1 else 1)
        if plain.upper() == "FREDERICK COUNTY, MARYLAND":
            self._county_title(text)
        elif level == 1 and plain.isupper() and len(plain) < 48:
            self.set_text_color(0, 0, 0)
            self.set_font("Inter", "B", 13)
            self.multi_cell(0, 6.4, text, align="C", markdown=True, new_x="LMARGIN", new_y="NEXT")
            self.ln(1.4)
        elif level == 2:
            self.set_text_color(*SECTION_BLUE)
            self.set_font("Inter", "B", 13)
            self.multi_cell(0, 6.2, text, align="L", markdown=True, new_x="LMARGIN", new_y="NEXT")
            self.ln(1)
        else:
            self.set_text_color(*NAVY if level >= 3 else (0, 0, 0))
            self.set_font("Inter", "B", 11 if level >= 3 else 15)
            self.multi_cell(0, 5.8, text, align="L", markdown=True, new_x="LMARGIN", new_y="NEXT")
            self.ln(0.8)
        self.set_text_color(*INK)
        self.set_font("Inter", "", 10.5)

    def _county_title(self, text: str) -> None:
        self.set_text_color(*NAVY)
        self.set_font("Inter", "B", 13)
        self.multi_cell(0, 6.6, text, align="C", markdown=True, new_x="LMARGIN", new_y="NEXT")
        y = self.get_y() + 0.8
        self.set_draw_color(*SECTION_BLUE)
        self.set_line_width(0.5)
        inset = 22
        self.line(self.l_margin + inset, y, self.w - self.r_margin - inset, y)
        self.ln(3.2)
        self.set_text_color(*INK)

    def rich(self, text: str, size: float = 10.5, h: float = 5.2) -> None:
        plain = re.sub(r"[*_]", "", text).strip()
        if plain.upper() == "FREDERICK COUNTY, MARYLAND":
            self._county_title(asterisks_to_fpdf(text))
            return
        self.set_text_color(*INK)
        self.set_font("Inter", "", size)
        rendered = asterisks_to_fpdf(text)
        # A URL on the same line as its label otherwise justifies the label
        # across the full measure.
        align = "L" if "://" in rendered else "J"
        self.multi_cell(
            0, h, rendered, align=align, markdown=True, new_x="LMARGIN", new_y="NEXT"
        )
        self.ln(1.4)

    def bullet(self, text: str) -> None:
        self.set_text_color(*INK)
        self.set_font("Inter", "", 10.5)
        self.multi_cell(
            0,
            5.2,
            "•  " + asterisks_to_fpdf(text),
            align="L",
            markdown=True,
            new_x="LMARGIN",
            new_y="NEXT",
        )
        self.ln(0.6)

    def render_table(self, rows: list[list[str]]) -> None:
        if not rows:
            return
        usable = self.w - self.l_margin - self.r_margin
        cols = max(len(r) for r in rows)
        rows = [r + [""] * (cols - len(r)) for r in rows]
        if cols == 2:
            weights = (0.22, 0.78)
        elif cols == 3:
            weights = (0.34, 0.40, 0.26)
        else:
            weights = tuple(1 / cols for _ in range(cols))
        col_widths = tuple(usable * w for w in weights)
        self.set_font("Inter", "", 9)
        with super().table(
            col_widths=col_widths,
            text_align="LEFT",
            line_height=4.6,
            markdown=True,
            first_row_as_headings=True,
            padding=1.2,
        ) as tbl:
            for i, row in enumerate(rows):
                pdf_row = tbl.row()
                for cell in row:
                    pdf_row.cell(asterisks_to_fpdf(cell), align="L")
                if i == 0:
                    pass
        self.ln(2)


def render_markdown(pdf: PacketPDF, text: str) -> None:
    text = re.sub(r"<a\s+[^>]*>|</a>", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    para: list[str] = []

    def flush() -> None:
        if not para:
            return
        pdf.rich(" ".join(s.strip() for s in para))
        para.clear()

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("<a ") or stripped.startswith("</a"):
            i += 1
            continue

        if (
            stripped.startswith("|")
            and i + 1 < len(lines)
            and re.match(r"^\s*\|(?:\s*:?-{1,}:?\s*\|)+\s*$", lines[i + 1])
        ):
            flush()
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|(?:\s*:?-{1,}:?\s*\|)+\s*$", lines[i]):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    rows.append(cells)
                i += 1
            pdf.render_table(rows)
            continue

        if not stripped or stripped == "---":
            flush()
            i += 1
            continue

        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            pdf.heading(stripped[level:].strip(), level)
            i += 1
            continue

        if stripped.startswith("- "):
            flush()
            pdf.bullet(stripped[2:].strip())
            i += 1
            continue

        para.append(stripped)
        i += 1

    flush()


def write_pdf(path: Path, markdown: str) -> None:
    pdf = PacketPDF()
    pdf.add_page()
    render_markdown(pdf, markdown)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    print(f"wrote {path.relative_to(ROOT)} ({path.stat().st_size} bytes, {pdf.page_no()} pages)")


def write_docx(path: Path, markdown: str) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    for level, size in ((1, 16), (2, 13), (3, 12)):
        heading_style = doc.styles[f"Heading {level}"]
        heading_style.font.color.rgb = RGBColor(0x17, 0x36, 0x5D)
        heading_style.font.size = Pt(size)
    sep = re.compile(r"^\s*\|(?:\s*:?-{1,}:?\s*\|)+\s*$")
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line == "---" or line.startswith("<"):
            i += 1
            continue
        if (
            line.startswith("|")
            and i + 1 < len(lines)
            and sep.match(lines[i + 1])
        ):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not sep.match(lines[i]):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    rows.append(cells)
                i += 1
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    cell_text = row[ci] if ci < len(row) else ""
                    paragraph = table.cell(ri, ci).paragraphs[0]
                    add_md_runs(paragraph, cell_text, force_bold=(ri == 0))
            continue
        if line.startswith("#"):
            level = min(3, len(line) - len(line.lstrip("#")))
            text = line.lstrip("#").strip()
            heading = doc.add_heading("", level=level)
            for run in list(heading.runs):
                run._element.getparent().remove(run._element)
            add_md_runs(heading, text)
            i += 1
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_md_runs(paragraph, line[2:])
            i += 1
            continue
        paragraph = doc.add_paragraph()
        add_md_runs(paragraph, line)
        i += 1
    doc.save(path)
    print(f"wrote {path.relative_to(ROOT)}")


def main() -> None:
    memo = "\n".join(
        [
            (SRC / "memo-preamble.md").read_text(encoding="utf-8").strip(),
            slice_between(PACKET, "## 1. Request", "# Grandfather Clause"),
            (SRC / "memo-sources.md").read_text(encoding="utf-8").strip(),
        ]
    )
    handout = (
        "STATEMENT AND ONE-PAGE HANDOUT\n\n"
        "Read at the mic. Leave the handout on the table.\n\n"
        + slice_between(PACKET, "## Resident handout / script", "## Resident proposal letter")
    )
    letter = slice_between(PACKET, "## Resident proposal letter", "## Sources")

    write_pdf(DOCS / "DABonData.pdf", PACKET)
    write_docx(DOCS / "DABonData.docx", PACKET)
    write_pdf(DOCS / "DAB_Energy_Trust_Memo.pdf", memo)
    write_pdf(DOCS / "DAB_Energy_Trust_Bylaws.pdf", (SRC / "bylaws.md").read_text(encoding="utf-8"))
    write_pdf(DOCS / "DAB_Script_Handout.pdf", handout)
    write_pdf(DOCS / "DAB_Resident_Proposal_Letter.pdf", letter)
    write_pdf(DOCS / "DAB_Example_Email.pdf", (SRC / "email.md").read_text(encoding="utf-8"))


if __name__ == "__main__":
    main()
